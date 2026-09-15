"""Extract the supplied log reference workbook into a small JSON knowledge base.

Usage: python scripts/build_log_knowledge.py SOURCE.xlsx
Requires openpyxl only for conversion; consuming the JSON uses the standard library.
"""

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re

import openpyxl


SHEETS = {
    "시스템로그": ("system", 3),
    "시스템로그 개발자용": ("system_developer", 1),
    "알람로그": ("alarm", 3),
}


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def build(source):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    entries = []
    notes = {}
    try:
        for sheet_name, (log_type, first_row) in SHEETS.items():
            sheet = workbook[sheet_name]
            if first_row > 1:
                notes[sheet_name] = sheet.cell(1, 1).value
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                values = [cell.value for cell in cells]
                if row_number < first_row or not any(v is not None for v in values):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Formula requires review: {sheet_name}:{row_number}")
                if any(v is not None for v in values[7:]):
                    raise ValueError(f"Unexpected columns: {sheet_name}:{row_number}")
                index, category, template, code, description, sixth, seventh = values[:7]
                if index is None or template is None:
                    raise ValueError(f"Incomplete record: {sheet_name}:{row_number}")
                prefix = re.match(r"^(?:\[[^\]\r\n]+\])+", template)
                entries.append({
                    "id": f"{log_type}:{text(index)}",
                    "log_type": log_type,
                    "source_index": text(index),
                    "category": category,
                    "code": text(code),
                    "message_template": template,
                    "tags": re.findall(r"\[([^\]]+)\]", prefix.group()) if prefix else [],
                    "description": description,
                    "example": seventh if log_type == "alarm" else sixth,
                    "recommended_action": sixth if log_type == "alarm" else None,
                    "source_note": None if log_type == "alarm" else seventh,
                    "reference_only": category == "기존 미출력 / 중복",
                    "source": {"sheet": sheet_name, "row": row_number},
                })
    finally:
        workbook.close()
    if len({entry["id"] for entry in entries}) != len(entries):
        raise ValueError("Duplicate IDs within a log type")
    return {
        "schema_version": 1,
        "source_file": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_notes": notes,
        "usage_notes": [
            "로그 종류와 코드로 후보를 찾고 메시지 문구를 함께 비교한다. 코드만으로 의미를 확정하지 않는다.",
            "category와 code는 심각도를 뜻하지 않는다. 원본에 없는 원인이나 해결방법은 추정하여 채우지 않았다.",
            "message_template은 원본 문구이며 정규식이 아니다. printf 자리표시자와 START/END 등의 설명 표기가 포함될 수 있다.",
            "reference_only 항목은 원본에서 기존 미출력 / 중복으로 분류한 참고 자료다.",
            "example에 구버전 또는 언어별 문구가 있을 수 있으므로 실제 로그와 함께 확인한다.",
            "source_note의 코드 근거는 제공 문서의 기재 내용이며 실제 소스 코드와 대조 검증하지 않았다.",
        ],
        "counts": dict(Counter(entry["log_type"] for entry in entries)),
        "entries": entries,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path,
                        default=Path(__file__).resolve().parents[1] / "knowledge" / "log_reference.json")
    args = parser.parse_args()
    result = build(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"counts": result["counts"], "reference_only": sum(e["reference_only"] for e in result["entries"])}, ensure_ascii=False))
